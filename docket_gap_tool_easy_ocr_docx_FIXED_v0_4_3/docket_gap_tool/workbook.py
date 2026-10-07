from __future__ import annotations

from pathlib import Path
import re
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .util import safe_sheet_name

HEADERS = [
    "Case",
    "Entry #",
    "Date",
    "Bloomberg Status",
    "Action",
    "Theme / Category",
    "Description",
]


def _style_sheet(ws, rows_count: int, table_name: str | None = None):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        1: 26,
        2: 10,
        3: 14,
        4: 18,
        5: 18,
        6: 34,
        7: 90,
    }
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if table_name and rows_count >= 1:
        end_row = max(rows_count, 1)
        end_col = len(HEADERS)
        ref = f"A1:{get_column_letter(end_col)}{end_row}"
        tab = Table(displayName=table_name, ref=ref)
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)


def _append_rows(ws, rows: list[dict]):
    ws.append(HEADERS)
    for row in rows:
        ws.append([
            row.get("case"),
            row.get("entry_number"),
            row.get("date"),
            row.get("bloomberg_status"),
            row.get("action"),
            row.get("theme"),
            row.get("description"),
        ])


def make_output_workbook(path: str | Path, case_rows: dict[str, list[dict]], summaries: list[dict], audit_notes: list[str] | None = None) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    used: set[str] = set()

    # Summary sheet
    ws = wb.create_sheet("Summary")
    summary_headers = [
        "Case", "Final needed", "Downloadable now", "Request needed", "Review optional", "Docket rows", "Unique downloaded entries", "Missing before skips"
    ]
    ws.append(summary_headers)
    for s in summaries:
        ws.append([
            s.get("case"), s.get("final_keep_request"), s.get("downloadable_now"), s.get("request_needed"),
            s.get("review_optional"), s.get("docket_rows_parsed"), s.get("unique_downloaded_entries"), s.get("missing_before_skips")
        ])
    total_row = ws.max_row + 1
    ws.append(["TOTAL", f"=SUM(B2:B{total_row-1})", f"=SUM(C2:C{total_row-1})", f"=SUM(D2:D{total_row-1})", f"=SUM(E2:E{total_row-1})", f"=SUM(F2:F{total_row-1})", f"=SUM(G2:G{total_row-1})", f"=SUM(H2:H{total_row-1})"])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(summary_headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.freeze_panes = "A2"

    # Audit notes
    ws = wb.create_sheet("Audit Notes")
    ws.append(["Note"])
    notes = audit_notes or []
    for n in notes:
        ws.append([n])
    ws.column_dimensions["A"].width = 120
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].font = Font(color="FFFFFF", bold=True)

    all_rows: list[dict] = []
    for rows in case_rows.values():
        all_rows.extend(rows)

    for name, rows in [
        ("All Final", all_rows),
        ("Downloadable Now", [r for r in all_rows if r.get("action") == "Downloadable now"]),
        ("Request Needed", [r for r in all_rows if r.get("action") == "Request needed"]),
    ]:
        ws = wb.create_sheet(safe_sheet_name(name, used))
        _append_rows(ws, rows)
        _style_sheet(ws, ws.max_row, re.sub(r"[^A-Za-z0-9]", "", name)[:20] + "Table")

    for case_name, rows in case_rows.items():
        ws = wb.create_sheet(safe_sheet_name(case_name, used))
        _append_rows(ws, rows)
        table_name = re.sub(r"[^A-Za-z0-9]", "", case_name)[:20] + "Table"
        if not table_name or table_name[0].isdigit():
            table_name = "Case" + table_name
        _style_sheet(ws, ws.max_row, table_name)

    wb.save(path)
